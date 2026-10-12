"""Write the synthetic Tier 2 intake fixtures to tests/fixtures/screening/. Deterministic; safe to rerun.

  python scripts/make_screening_fixtures.py
"""
import csv
import random
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tests"))
import screening_fixture_data as F  # noqa: E402

OUT = ROOT / "tests" / "fixtures" / "screening"
RR_HEADERS = ["Unit", "Type", "Bd/Ba", "Sq Ft", "Status", "Resident", "Phone", "Lease From", "Lease To",
              "Rent", "Market", "Concession", "Deposit"]


# ---------------------------------------------------------------- rent roll rows

def rr_rows(units: list[dict]) -> list[list]:
    rows = []
    for i, u in enumerate(units):
        occupied = u["status"] == "Occupied"
        rows.append([
            u["unit"], u["type"], u["bdba"], u["sqft"], u["status"],
            F.TENANTS[i] if occupied else "", F.PHONES[i] if occupied else "",
            u["start"], "MTM" if occupied and u["end"] is None else u["end"],
            u["rent"], u["market"], u["concession"] or None, u["deposit"],
        ])
    return rows


def rr_totals(units: list[dict]) -> list:
    return ["Totals", "", "", sum(u["sqft"] for u in units), "", "", "", "", "",
            sum(u["rent"] or 0 for u in units), sum(u["market"] for u in units), None, None]


def fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, date):
        return v.strftime("%m/%d/%Y")
    if isinstance(v, float) or (isinstance(v, int) and not isinstance(v, bool) and v >= 1000):
        return f"{v:,.2f}" if isinstance(v, float) else f"{v:,}"
    return str(v)


# ---------------------------------------------------------------- T-12 rows

def t12_rows(spec: dict, problem: bool) -> list[tuple[str, list | None, float | None, str]]:
    """(label, monthly values or None for a section header, total, kind)."""
    months = F.months_for(problem)
    rows: list = []
    sums = {}
    for section, lines in spec.items():
        rows.append((section, None, None, "section"))
        sec_month = [0.0] * len(months)
        for label, amount in lines:
            vals = [round(float(F.monthly(amount, m)), 2) for m in months]
            sec_month = [a + b for a, b in zip(sec_month, vals)]
            rows.append((label, vals, round(sum(vals), 2), "line"))
        sums[section] = [round(v, 2) for v in sec_month]
        if section == "INCOME":
            rows.append(("Total Income", sums[section], round(sum(sums[section]), 2), "subtotal"))
        elif section == "OPERATING EXPENSES":
            total = round(sum(sums[section]), 2) + (F.TOTAL_EXPENSES_PLANTED_ERROR if problem else 0)
            rows.append(("Total Operating Expenses", sums[section], total, "subtotal"))
            noi = [round(a - b, 2) for a, b in zip(sums["INCOME"], sums[section])]
            rows.append(("Net Operating Income", noi, round(sum(noi), 2), "subtotal"))
    return rows


def money(v: float | None) -> str:
    if v is None:
        return ""
    return f"({abs(v):,.2f})" if v < 0 else f"{v:,.2f}"


# ---------------------------------------------------------------- writers

def write_rr_xlsx(path: Path, units: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rent Roll"
    ws.append([F.PROPERTY])
    ws.append(["Rent Roll as of 09/30/2026"])
    ws.append([])
    ws.append(RR_HEADERS)
    for row in rr_rows(units):
        ws.append(row)
        for c in (8, 9):
            if isinstance(ws.cell(ws.max_row, c).value, date):
                ws.cell(ws.max_row, c).number_format = "mm/dd/yyyy"
    ws.append(rr_totals(units))
    wb.save(path)


def write_rr_csv(path: Path, units: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(RR_HEADERS)
        for row in rr_rows(units):
            w.writerow([fmt_cell(v) if isinstance(v, date) else ("" if v is None else v) for v in row])


def write_t12_xlsx(path: Path, spec: dict, problem: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "T12"
    ws.append([F.PROPERTY])
    ws.append(["Trailing 12 Month Operating Statement"])
    ws.append([])
    ws.append(["Account"] + [F.MONTH_LABELS[m] for m in F.months_for(problem)] + ["Total"])
    for label, vals, total, kind in t12_rows(spec, problem):
        ws.append([label] if vals is None else [label] + vals + [total])
    wb.save(path)


def write_t12_csv(path: Path, spec: dict, problem: bool) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Account"] + [F.MONTH_LABELS[m] for m in F.months_for(problem)] + ["Total"])
        for label, vals, total, _ in t12_rows(spec, problem):
            w.writerow([label] if vals is None else [label] + [f"{v:.2f}" for v in vals] + [f"{total:.2f}"])


def _pdf_table(path: Path, title: str, headers: list[str], rows: list[list[str]], widths: list[float],
               right: set[int], font_size: float) -> None:
    """A plain text table with no grid lines, like most accounting exports."""
    c = canvas.Canvas(str(path), pagesize=landscape(letter))
    c.setTitle(title)
    width, height = landscape(letter)
    xs = [36]
    for w in widths[:-1]:
        xs.append(xs[-1] + w)
    y = height - 40
    c.setFont("Helvetica-Bold", 11)
    c.drawString(36, y, F.PROPERTY)
    c.setFont("Helvetica", 9)
    c.drawString(36, y - 14, title)
    y -= 40

    def draw_row(cells, bold=False):
        nonlocal y
        if y < 40:
            c.showPage()
            y = height - 40
        c.setFont("Helvetica-Bold" if bold else "Helvetica", font_size)
        for i, text in enumerate(cells):
            if i in right:
                c.drawRightString(xs[i] + widths[i] - 4, y, text)
            else:
                c.drawString(xs[i], y, text)
        y -= font_size + 5

    draw_row(headers, bold=True)
    for r in rows:
        draw_row(r, bold=r and r[0].isupper() and all(not x for x in r[1:]))
    c.save()


def write_rr_pdf(path: Path, units: list[dict]) -> None:
    rows = [[fmt_cell(v) for v in r] for r in rr_rows(units)] + [[fmt_cell(v) for v in rr_totals(units)]]
    _pdf_table(path, "Rent Roll as of 09/30/2026", RR_HEADERS, rows,
               [40, 55, 38, 42, 52, 95, 58, 62, 62, 56, 56, 58, 50], {3, 9, 10, 11, 12}, 7.5)


def write_t12_pdf(path: Path, spec: dict, problem: bool) -> None:
    months = F.months_for(problem)
    headers = ["Account"] + [F.MONTH_LABELS[m] for m in months] + ["Total"]
    rows = [[label] if vals is None else [label] + [money(v) for v in vals] + [money(total)]
            for label, vals, total, _ in t12_rows(spec, problem)]
    widths = [128] + [44] * len(months) + [56]
    _pdf_table(path, "Trailing 12 Month Operating Statement", headers, rows, widths,
               set(range(1, len(headers))), 6.3)


def write_rr_scanned(path: Path, units: list[dict]) -> None:
    """Rent roll as an image-only PDF (no text layer), like a scanned printout."""
    rng = random.Random(7)
    img = Image.new("L", (2200, 1700), 255)
    d = ImageDraw.Draw(img)
    fonts = Path("C:/Windows/Fonts")
    reg = ImageFont.truetype(str(fonts / "arial.ttf"), 24)
    bold = ImageFont.truetype(str(fonts / "arialbd.ttf"), 24)
    d.text((70, 60), F.PROPERTY, font=ImageFont.truetype(str(fonts / "arialbd.ttf"), 32), fill=0)
    d.text((70, 105), "Rent Roll as of 09/30/2026", font=reg, fill=0)
    cols = [70, 165, 300, 395, 495, 625, 870, 1020, 1185, 1350, 1490, 1630, 1790]
    right_edge = {3: 470, 9: 1460, 10: 1600, 11: 1760, 12: 1910}
    y = 180

    def row(cells, font):
        nonlocal y
        for i, text in enumerate(cells):
            if i in right_edge and text:
                w = d.textlength(text, font=font)
                d.text((right_edge[i] - w, y), text, font=font, fill=0)
            else:
                d.text((cols[i], y), text, font=font, fill=0)
        y += 52

    row(RR_HEADERS, bold)
    for r in rr_rows(units) + [rr_totals(units)]:
        row([fmt_cell(v) for v in r], reg)
    # Scanner look: a little blur and speckle, no rotation.
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    px = img.load()
    for _ in range(9000):
        x, yy = rng.randrange(img.width), rng.randrange(img.height)
        px[x, yy] = rng.randrange(150, 230)
    img.save(path, "PDF", resolution=200)


def write_failures() -> None:
    (OUT / "bad_not_really_excel.xlsx").write_bytes(b"This is plain text pretending to be a spreadsheet.\n" * 20)
    wb = Workbook()
    wb.active.append([F.PROPERTY])
    wb.active.append(["(no data)"])
    wb.save(OUT / "empty_workbook.xlsx")
    c = canvas.Canvas(str(OUT / "password_protected.pdf"), encrypt="fixture-secret")
    c.drawString(72, 720, "Protected rent roll")
    c.save()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_rr_xlsx(OUT / "problem_rent_roll.xlsx", F.PROBLEM_UNITS)
    write_rr_pdf(OUT / "problem_rent_roll.pdf", F.PROBLEM_UNITS)
    write_rr_scanned(OUT / "problem_rent_roll_scanned.pdf", F.PROBLEM_UNITS)
    write_t12_xlsx(OUT / "problem_t12.xlsx", F.PROBLEM_T12, problem=True)
    write_t12_pdf(OUT / "problem_t12.pdf", F.PROBLEM_T12, problem=True)
    write_rr_xlsx(OUT / "clean_rent_roll.xlsx", F.CLEAN_UNITS)
    write_rr_csv(OUT / "clean_rent_roll.csv", F.CLEAN_UNITS)
    write_t12_xlsx(OUT / "clean_t12.xlsx", F.CLEAN_T12, problem=False)
    write_t12_csv(OUT / "clean_t12.csv", F.CLEAN_T12, problem=False)
    write_failures()
    for p in sorted(OUT.iterdir()):
        print(f"{p.name:34} {p.stat().st_size:>8,} bytes")


if __name__ == "__main__":
    main()
