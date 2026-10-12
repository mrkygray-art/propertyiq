"""Milestone 9: deal PDF and Excel exports match the app; Excel formulas are live (checked in OpenOffice)."""
import json
import os
import re
import subprocess
from io import BytesIO
from pathlib import Path

import psycopg
import pytest
from openpyxl import load_workbook

import screening_fixture_data as F
from propertyiq.config import get_settings
from propertyiq.llm import LLMResult
from test_screening_exports import OO_PYTHON, RECALC, pdf_text

pytestmark = pytest.mark.db

try:
    with psycopg.connect(get_settings().dsn("piq_readonly")) as _c:
        _c.execute("SELECT 1 FROM screening_adjustments LIMIT 1")
except Exception as exc:  # pragma: no cover
    pytest.skip(f"Tier 2 tables not set up: {exc}", allow_module_level=True)

from fastapi.testclient import TestClient  # noqa: E402

from propertyiq import main  # noqa: E402
from propertyiq.screening import compare  # noqa: E402
from propertyiq.screening.export_deal import DISCLAIMER  # noqa: E402

FX = Path(__file__).parent / "fixtures" / "screening"
LOAN = {"value": 6000000, "loan_amount": 4000000, "interest_rate_pct": 6.5, "amortization_years": 30,
        "interest_only_years": 0, "term_years": 10}
ROW = {k: i for i, k in enumerate((k for k, *_ in compare.LINES), start=5)}   # Underwritten sheet rows
T0 = 5 + len(compare.LINES) + 1                                               # first totals row
NOI_ROW, DSCR_ROW = T0 + 5, T0 + 7


class FakeAI:
    name = "fake"

    def available(self):
        return True

    def chat(self, system, user, **_):
        return LLMResult('{"1": "contract_services"}', "fake", None, None, None)


@pytest.fixture(scope="module")
def setup():
    main.PROVIDER, saved = FakeAI(), main.PROVIDER
    with TestClient(main.app) as client:
        r = client.post("/api/auth/login", json={"username": "analyst", "password": os.environ["PIQ_DEMO_PASSWORD"]})
        h = {"Authorization": f"Bearer {r.json()['token']}"}
        did = client.post("/api/screening/deals", headers=h, json={
            "name": "pytest exports", "property_facts": {**F.DEAL_FACTS, "property_type": "multifamily"},
            "borrower_inputs": {"property_type": "multifamily"}}).json()["deal_id"]
        try:
            for name in ("problem_rent_roll.xlsx", "problem_t12.xlsx"):
                doc = client.post(f"/api/screening/deals/{did}/documents", headers=h,
                                  files={"file": (name, (FX / name).read_bytes())}).json()
                client.post(f"/api/screening/documents/{doc['document_id']}/approve", headers=h)
            client.get(f"/api/screening/deals/{did}/analysis", headers=h)
            client.put(f"/api/screening/deals/{did}/california", headers=h, json={"facts": {}, "condition": [
                {"item_key": "roof_flat", "details": {"age_years": "21"}, "estimated_cost": 70000}]})
            for kind in ("rent", "sale"):
                client.post(f"/api/screening/deals/{did}/market", headers=h, json={
                    "kind": kind, "address": "1200 Juniper Ct", "city": "Los Angeles", "radius_miles": 1, "max_age_months": 12})
            client.put(f"/api/screening/deals/{did}/loan", headers=h, json=LOAN)
            s = client.post(f"/api/screening/deals/{did}/adjustments/suggest", headers=h).json()
            for a in s["adjustments"]:
                client.patch(f"/api/screening/adjustments/{a['adjustment_id']}", headers=h, json={"status": "approved"})
            comp = client.get(f"/api/screening/deals/{did}/compare", headers=h).json()
            pdf = client.get(f"/api/screening/deals/{did}/export/pdf", headers=h)
            xlsx = client.get(f"/api/screening/deals/{did}/export/xlsx", headers=h)
            yield {"compare": comp, "pdf": pdf, "xlsx": xlsx}
        finally:
            client.delete(f"/api/screening/deals/{did}", headers=h)
            main.PROVIDER = saved


def test_pdf_contents(setup):
    pdf = setup["pdf"]
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    pages = len(re.findall(rb"/Type\s*/Page(?!s)", pdf.content))
    text = pdf_text(pdf.content)
    assert pages >= 3 and text.count(DISCLAIMER) >= pages   # disclaimer on every page
    c = setup["compare"]
    # The summary sentence wraps across lines in the PDF, so check its opening words.
    for expected in (c["summary"][:40], "Borrower vs underwritten", "TAX_REASSESSMENT", "Prop 13", "Market support",
                     "SYNTHETIC DEMO DATA", "Data sources", "UNVERIFIED", "problem_t12.xlsx"):
        assert expected in text, expected
    for who in F.TENANTS + F.PHONES:
        assert who not in text


def test_workbook_structure(setup):
    wb = load_workbook(BytesIO(setup["xlsx"].content))
    assert {"Borrower", "Adjustments", "Underwritten", "Flags", "Comps", "Inputs", "Calculations"} <= set(wb.sheetnames)
    assert str(wb["Underwritten"][f"C{ROW['real_estate_taxes']}"].value).startswith("=IF(COUNTIFS(")
    assert str(wb["Inputs"]["B8"].value).startswith("=Underwritten!")   # PGI comes from the Underwritten sheet
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    assert not any(who in text for who in F.TENANTS + F.PHONES)


def _recalc(tmp_path, data: bytes) -> dict:
    if not OO_PYTHON.exists():
        pytest.skip("OpenOffice not installed")
    src, dst = tmp_path / "deal.xlsx", tmp_path / "deal.json"
    src.write_bytes(data)
    subprocess.run([str(OO_PYTHON), str(RECALC), str(src), str(dst)], check=True, timeout=180)
    return json.loads(dst.read_text())


def test_excel_recalculates_to_the_app_numbers(setup, tmp_path):
    cells = _recalc(tmp_path, setup["xlsx"].content)
    c = setup["compare"]
    u, b = cells["Underwritten"], c["totals"]
    assert u[f"C{NOI_ROW}"] == pytest.approx(b["underwritten"]["noi"], abs=0.01)
    assert u[f"D{NOI_ROW}"] == pytest.approx(b["borrower"]["noi"], abs=0.01)
    assert u[f"C{DSCR_ROW}"] == pytest.approx(c["underwritten_case"]["metrics"]["dscr"], abs=1e-6)
    assert cells["Calculations"]["B6"] == pytest.approx(c["underwritten_case"]["metrics"]["noi"], abs=0.01)
    assert cells["Calculations"]["B10"] == pytest.approx(c["underwritten_case"]["metrics"]["dscr"], abs=1e-6)
    assert cells["Guidelines"]["E14"] == c["underwritten_case"]["overall"]
    assert not any(isinstance(v, dict) for sheet in cells.values() for v in sheet.values()), "formula errors"


def test_excel_adjustments_are_live(setup, tmp_path):
    """Reject the tax adjustment inside the workbook: NOI should rise by exactly the tax difference."""
    wb = load_workbook(BytesIO(setup["xlsx"].content))
    wa = wb["Adjustments"]
    tax_row = next(r for r in range(4, wa.max_row + 1) if wa.cell(r, 2).value == "real_estate_taxes")
    uw_tax, b_tax = wa.cell(tax_row, 5).value, wa.cell(tax_row, 4).value
    wa.cell(tax_row, 6).value = "rejected"
    buf = BytesIO()
    wb.save(buf)
    cells = _recalc(tmp_path, buf.getvalue())
    before = setup["compare"]["totals"]["underwritten"]["noi"]
    assert cells["Underwritten"][f"C{NOI_ROW}"] == pytest.approx(before + (uw_tax - b_tax), abs=0.01)
    assert cells["Underwritten"][f"C{ROW['real_estate_taxes']}"] == pytest.approx(b_tax, abs=0.01)
