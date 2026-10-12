"""Excel workbook with live formulas (openpyxl).

Sheets: Inputs (editable), Calculations, Guidelines, Stress (scenarios + DSCR grid).
Every result is an Excel formula that points back at the Inputs sheet, so a lender can
change an input and watch everything update. The formulas mirror calc.py; a test opens
the file in OpenOffice, lets it recalculate, and checks the results match the engine.
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from propertyiq.screening import DISCLAIMER
from propertyiq.screening.report import TYPE_LABELS, Report

USD = '"$"#,##0;[Red]-"$"#,##0'
USD_CENTS = '"$"#,##0.00;[Red]-"$"#,##0.00'
RATIO = "0.00"
PCT1 = "0.0%"
PCT2 = "0.00%"
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
NOTE = Font(italic=True, color="666666", size=9)
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")  # editable cells
RED_FILL = PatternFill("solid", fgColor="F8D7D7")
GREEN_FONT = Font(color="0A7D0A", bold=True)
RED_FONT = Font(color="C03030", bold=True)
GRAY_FONT = Font(color="6B6A66", bold=True)


def _f(v: Decimal | int | None):
    return None if v is None else float(v)


def _status_formatting(ws, rng: str) -> None:
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], font=GREEN_FONT))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], font=RED_FONT))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"NOT CHECKED"'], font=GRAY_FONT))


def render_xlsx(r: Report) -> bytes:
    assert r.valid
    d, g = r.deal, r.guidelines
    wb = Workbook()

    # ------------------------------------------------------------ Inputs
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = f"Loan screening inputs: {d.property_name or '(unnamed property)'}"
    ws["A1"].font = TITLE
    ws["A2"] = DISCLAIMER
    ws["A2"].font = NOTE
    ws["A3"] = "Yellow cells are inputs. Change them and every other sheet recalculates."
    ws["A3"].font = NOTE
    rows = [
        ("name", "Property", d.property_name, None, ""),
        ("type", "Property type", d.property_type, None, "multifamily, office, retail, industrial, mixed_use, other"),
        ("value", "Value used for screening", _f(d.value), USD, "Purchase price or appraised value"),
        ("pgi", "Potential gross income (PGI)", _f(d.pgi), USD, "Rent if 100% full, per year"),
        ("other", "Other income", _f(d.other_income), USD, "Parking, laundry, fees, per year"),
        ("vac", "Vacancy and credit loss", _f(d.vacancy), PCT2, "Share of income lost to empty units and non-payment"),
        ("opex", "Operating expenses", _f(d.operating_expenses), USD, "Taxes, insurance, utilities, repairs, management, payroll"),
        ("res", "Replacement reserves", _f(d.reserves), USD, "Yearly set-aside for big repairs"),
        ("loan", "Loan amount", _f(d.loan_amount) or 0, USD, ""),
        ("rate", "Interest rate (per year)", _f(d.interest_rate), PCT2, ""),
        ("amort", "Amortization (years)", d.amortization_years, "0", "Years over which the loan is paid off"),
        ("io", "Interest-only period (years)", _f(d.interest_only_years), "0.0", "0 if none"),
        ("term", "Loan term (years)", _f(d.term_years), "0.0", "Informational"),
    ]
    ref: dict[str, str] = {}
    for i, (key, label, value, fmt, note) in enumerate(rows, start=5):
        ws.cell(i, 1, label)
        c = ws.cell(i, 2, value)
        c.fill = INPUT_FILL
        if fmt:
            c.number_format = fmt
        ws.cell(i, 3, note).font = NOTE
        ref[key] = f"Inputs!$B${i}"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 60

    # ------------------------------------------------------------ Calculations
    wc = wb.create_sheet("Calculations")
    wc["A1"] = "Calculations (live formulas)"
    wc["A1"].font = TITLE
    wc["A2"] = DISCLAIMER
    wc["A2"].font = NOTE
    for col, head in enumerate(("Metric", "Value", "Formula in plain words"), start=1):
        wc.cell(4, col, head).font = BOLD
    L, V, P, O, R, RATE, N, IO = ref["loan"], ref["value"], ref["pgi"], ref["other"], ref["res"], ref["rate"], ref["amort"], ref["io"]
    calc_rows = [
        ("egi", "Effective gross income (EGI)", f"=({P}+{O})*(1-{ref['vac']})", USD, "(PGI + other income) x (1 - vacancy)"),
        ("noi", "Net operating income (NOI)", f"=B5-{ref['opex']}-{R}", USD, "EGI - operating expenses - reserves"),
        ("pay", "Monthly loan payment", f'=IF({L}<=0,"needs a loan amount",IF({RATE}=0,{L}/({N}*12),'
                                         f"{L}*({RATE}/12)/(1-(1+{RATE}/12)^(-{N}*12))))", USD_CENTS,
         "Loan x r / (1 - (1 + r)^-n); r = rate / 12, n = years x 12; Loan / n if the rate is 0"),
        ("ads", "Annual debt service", '=IF(ISNUMBER(B7),12*B7,"n/a")', USD, "12 x monthly payment"),
        ("iods", "Interest-only debt service", f'=IF(AND({IO}>0,{L}>0),{L}*{RATE},"n/a")', USD, "Loan x rate, during the interest-only period"),
        ("dscr", "DSCR", '=IF(AND(ISNUMBER(B8),B8<>0),B6/B8,"n/a")', RATIO, "NOI / annual debt service"),
        ("dscr_io", "DSCR (interest-only)", '=IF(AND(ISNUMBER(B9),B9<>0),B6/B9,"n/a")', RATIO, "NOI / interest-only debt service"),
        ("ltv", "Loan-to-value (LTV)", f'=IF({L}>0,{L}/{V},"n/a")', PCT1, "Loan / value"),
        ("dy", "Debt yield", f'=IF({L}>0,B6/{L},"n/a")', PCT1, "NOI / loan"),
        ("cap", "Cap rate (implied)", f"=B6/{V}", PCT2, "NOI / value"),
        ("be", "Break-even occupancy", f'=IF(AND(ISNUMBER(B8),({P}+{O})>0),({ref["opex"]}+{R}+B8)/({P}+{O}),"n/a")', PCT1,
         "(Operating expenses + reserves + annual debt service) / (PGI + other income)"),
        ("const", "Debt constant", f'=IF(AND({L}>0,ISNUMBER(B8)),B8/{L},"n/a")', PCT2, "Annual debt service / loan"),
    ]
    cref: dict[str, str] = {}
    for i, (key, label, formula, fmt, words) in enumerate(calc_rows, start=5):
        wc.cell(i, 1, label)
        c = wc.cell(i, 2, formula)
        c.number_format = fmt
        wc.cell(i, 3, words).font = NOTE
        cref[key] = f"Calculations!$B${i}"
    wc.column_dimensions["A"].width = 30
    wc.column_dimensions["B"].width = 18
    wc.column_dimensions["C"].width = 80

    # ------------------------------------------------------------ Guidelines
    wg = wb.create_sheet("Guidelines")
    wg["A1"] = f"Guidelines: {g.name}"
    wg["A1"].font = TITLE
    wg["A2"] = "Set your own limits. Example values are not industry standards."
    wg["A2"].font = NOTE
    wg["A3"] = "DSCR basis"
    wg["B3"] = g.dscr_basis
    wg["B3"].fill = INPUT_FILL
    wg["C3"] = "amortizing or interest_only (interest_only applies only when there is an interest-only period)"
    wg["C3"].font = NOTE
    for col, head in enumerate(("Rule", "Turned on?", "Limit", "Actual", "Result", "Check"), start=1):
        wg.cell(5, col, head).font = BOLD
    dscr_used = f'IF(AND($B$3="interest_only",ISNUMBER({cref["dscr_io"]})),{cref["dscr_io"]},{cref["dscr"]})'
    loan_or_na = f'IF({L}>0,{L},"n/a")'
    grules = [
        ("Minimum DSCR", g.min_dscr, RATIO, f"={dscr_used}", ">="),
        ("Maximum LTV", g.max_ltv, PCT1, f"={cref['ltv']}", "<="),
        ("Minimum debt yield", g.min_debt_yield, PCT1, f"={cref['dy']}", ">="),
        ("Maximum break-even occupancy", g.max_break_even, PCT1, f"={cref['be']}", "<="),
        ("Minimum loan size", g.min_loan, USD, f"={loan_or_na}", ">="),
        ("Maximum loan size", g.max_loan, USD, f"={loan_or_na}", "<="),
    ]
    gref: dict[str, str] = {}
    for i, (label, rule, fmt, actual, op) in enumerate(grules, start=6):
        wg.cell(i, 1, label)
        on = wg.cell(i, 2, "Yes" if rule.enabled else "No")
        on.fill = INPUT_FILL
        lim = wg.cell(i, 3, _f(rule.threshold))
        lim.fill = INPUT_FILL
        lim.number_format = fmt
        a = wg.cell(i, 4, actual)
        a.number_format = fmt
        wg.cell(i, 5, f'=IF(OR(B{i}<>"Yes",NOT(ISNUMBER(C{i})),NOT(ISNUMBER(D{i}))),"NOT CHECKED",'
                      f'IF(D{i}{op}C{i},"PASS","FAIL"))')
        wg.cell(i, 6, "minimum" if op == ">=" else "maximum").font = NOTE
        gref[label] = f"Guidelines!$C${i}"
    trow = 6 + len(grules)
    wg.cell(trow, 1, "Allowed property types")
    wg.cell(trow, 2, "Yes" if g.allowed_types_enabled else "No").fill = INPUT_FILL
    wg.cell(trow, 3, ",".join(g.allowed_types)).fill = INPUT_FILL
    wg.cell(trow, 4, f"={ref['type']}")
    wg.cell(trow, 5, f'=IF(B{trow}<>"Yes","NOT CHECKED",IF(ISNUMBER(SEARCH(","&D{trow}&",",","&C{trow}&",")),"PASS","FAIL"))')
    orow = trow + 2
    wg.cell(orow, 1, "Overall").font = BOLD
    wg.cell(orow, 5, f'=IF(COUNTIF(E6:E{trow},"FAIL")>0,"FAIL",IF(COUNTIF(E6:E{trow},"PASS")>0,"PASS","NOT CHECKED"))').font = BOLD
    _status_formatting(wg, f"E6:E{orow}")
    for col, w in zip("ABCDEF", (32, 12, 40, 16, 14, 10)):
        wg.column_dimensions[col].width = w
    loan_size_result = f'IF(OR(Guidelines!$E${6 + 4}="FAIL",Guidelines!$E${6 + 5}="FAIL",Guidelines!$E${trow}="FAIL"),"FAIL","")'

    # ------------------------------------------------------------ Stress
    wsx = wb.create_sheet("Stress")
    wsx["A1"] = "Stress tests (what if things get worse)"
    wsx["A1"].font = TITLE
    wsx["A2"] = "Yellow rows are the scenario changes; edit them to try your own."
    wsx["A2"].font = NOTE
    labels = [
        ("Interest rate change (points)", PCT2), ("Vacancy change (points)", PCT2), ("Collected income change", PCT1),
        ("Operating expense change", PCT1), ("Exit cap rate change (points)", PCT2),
        ("Interest rate", PCT2), ("Vacancy", PCT2), ("Effective gross income", USD), ("Operating expenses", USD),
        ("NOI", USD), ("Annual debt service", USD), ("DSCR", RATIO), ("Value used", USD), ("LTV", PCT1),
        ("Debt yield", PCT1), ("Break-even occupancy", PCT1), ("DSCR used for the check", RATIO),
        ("DSCR check", None), ("LTV check", None), ("Debt yield check", None), ("Break-even check", None), ("Result", None),
    ]
    wsx.cell(4, 1, "Scenario").font = BOLD
    for i, (label, _) in enumerate(labels, start=5):
        wsx.cell(i, 1, label)
    for j, s in enumerate(r.stress, start=2):
        col = get_column_letter(j)
        sc = s.scenario
        wsx.cell(4, j, sc.label).font = BOLD
        wsx.cell(4, j).alignment = Alignment(wrap_text=True, vertical="top")
        for i, v in enumerate((sc.rate_add, sc.vacancy_add, sc.income_change, sc.expense_change, sc.exit_cap_add), start=5):
            c = wsx.cell(i, j, float(v))
            c.fill = INPUT_FILL
        c_ = lambda row: f"{col}{row}"  # noqa: E731
        gross = f"({P}+{O})"
        formulas = {
            10: f"=MIN(1,{RATE}+{c_(5)})",
            11: f"=MIN(1,{ref['vac']}+{c_(6)})",
            12: f"={gross}*(1+{c_(7)})*(1-{c_(11)})",
            13: f"={ref['opex']}*(1+{c_(8)})",
            14: f"={c_(12)}-{c_(13)}-{R}",
            15: f'=IF({L}<=0,"n/a",12*IF({c_(10)}=0,{L}/({N}*12),{L}*({c_(10)}/12)/(1-(1+{c_(10)}/12)^(-{N}*12))))',
            16: f'=IF(AND(ISNUMBER({c_(15)}),{c_(15)}<>0),{c_(14)}/{c_(15)},"n/a")',
            17: f'=IF({c_(9)}=0,{V},IF(AND({c_(14)}>0,{cref["cap"]}+{c_(9)}>0),{c_(14)}/({cref["cap"]}+{c_(9)}),"n/a"))',
            18: f'=IF(AND({L}>0,ISNUMBER({c_(17)})),{L}/{c_(17)},"n/a")',
            19: f'=IF({L}>0,{c_(14)}/{L},"n/a")',
            20: f'=IF(AND(ISNUMBER({c_(15)}),{gross}*(1+{c_(7)})>0),({c_(13)}+{R}+{c_(15)})/({gross}*(1+{c_(7)})),"n/a")',
            21: (f'=IF(AND(Guidelines!$B$3="interest_only",{IO}>0,{L}>0,{L}*{c_(10)}<>0),{c_(14)}/({L}*{c_(10)}),{c_(16)})'),
        }
        for row, formula in formulas.items():
            wsx.cell(row, j, formula)
        checks = [(22, 21, "Minimum DSCR", ">="), (23, 18, "Maximum LTV", "<="), (24, 19, "Minimum debt yield", ">="),
                  (25, 20, "Maximum break-even occupancy", "<=")]
        for row, actual_row, rule_label, op in checks:
            on = gref[rule_label].replace("$C$", "$B$")
            lim = gref[rule_label]
            wsx.cell(row, j, f'=IF(OR({on}<>"Yes",NOT(ISNUMBER({lim})),NOT(ISNUMBER({c_(actual_row)}))),"NOT CHECKED",'
                             f'IF({c_(actual_row)}{op}{lim},"PASS","FAIL"))')
        base_rules = f'{loan_size_result}'
        wsx.cell(26, j, f'=IF(OR(COUNTIF({c_(22)}:{c_(25)},"FAIL")>0,{base_rules}="FAIL"),"FAIL",'
                        f'IF(OR(COUNTIF({c_(22)}:{c_(25)},"PASS")>0,'
                        f'COUNTIF(Guidelines!$E$10:$E${trow},"PASS")>0),"PASS","NOT CHECKED"))').font = BOLD
        wsx.column_dimensions[col].width = 16
    for i, (_, fmt) in enumerate(labels, start=5):
        if fmt:
            for j in range(2, 2 + len(r.stress)):
                wsx.cell(i, j).number_format = fmt
    last_col = get_column_letter(1 + len(r.stress))
    _status_formatting(wsx, f"B22:{last_col}26")
    wsx.column_dimensions["A"].width = 30
    wsx.row_dimensions[4].height = 42

    # Two-way DSCR grid: interest rate (rows) x vacancy (columns)
    top = 29
    wsx.cell(top, 1, "DSCR by interest rate (rows) and vacancy (columns)").font = BOLD
    wsx.cell(top + 1, 1, "Cells below the minimum DSCR are shaded red.").font = NOTE
    wsx.cell(top + 2, 1, "Rate \\ Vacancy").font = BOLD
    for j, vac in enumerate(r.grid_vacancies, start=2):
        c = wsx.cell(top + 2, j, float(vac))
        c.number_format = PCT1
        c.fill = INPUT_FILL
        c.font = BOLD
    for i, rate in enumerate(r.grid_rates, start=top + 3):
        rc = wsx.cell(i, 1, float(rate))
        rc.number_format = PCT2
        rc.fill = INPUT_FILL
        rc.font = BOLD
        for j in range(2, 2 + len(r.grid_vacancies)):
            col = get_column_letter(j)
            rate_c, vac_c = f"$A{i}", f"{col}${top + 2}"
            pay = f"IF({rate_c}=0,{L}/({N}*12),{L}*({rate_c}/12)/(1-(1+{rate_c}/12)^(-{N}*12)))"
            c = wsx.cell(i, j, f'=IF({L}<=0,"n/a",(({P}+{O})*(1-{vac_c})-{ref["opex"]}-{R})/(12*{pay}))')
            c.number_format = RATIO
    first = f"B{top + 3}"
    grid_rng = f"{first}:{get_column_letter(1 + len(r.grid_vacancies))}{top + 2 + len(r.grid_rates)}"
    wsx.conditional_formatting.add(grid_rng, FormulaRule(
        formula=[f'AND(ISNUMBER({first}),Guidelines!$B$6="Yes",ISNUMBER(Guidelines!$C$6),{first}<Guidelines!$C$6)'],
        fill=RED_FILL))

    for sheet in wb.worksheets:
        sheet.oddFooter.center.text = DISCLAIMER
        sheet.oddFooter.center.size = 8
    wb.properties.title = f"Loan screening - {d.property_name or 'deal'}"
    wb.properties.creator = "PropertyIQ"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
